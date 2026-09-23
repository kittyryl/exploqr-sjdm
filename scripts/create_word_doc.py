import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_document():
    doc = docx.Document()

    # Standard 1-inch margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Base Style: Plain Black Text
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    def add_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(18)
        run.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)

    def add_meta(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0, 0, 0)

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(14)
        run.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(12)
        run.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)

    def add_speech(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.left_indent = Inches(0.4)
        
        prefix = p.add_run("What to say: ")
        prefix.bold = True
        prefix.font.color.rgb = RGBColor(0, 0, 0)

        run = p.add_run(f'"{text}"')
        run.font.color.rgb = RGBColor(0, 0, 0)

    input_path = os.path.join(os.path.dirname(__file__), '..', 'DEVELOPMENT-PROCEDURE.md')
    input_path = os.path.abspath(input_path)

    with open(input_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue

        if stripped.startswith('# '):
            add_title(stripped[2:].strip())
        elif stripped.startswith('Research Team:') or stripped.startswith('Project:') or stripped.startswith('System:'):
            add_meta(stripped)
        elif stripped == '---':
            continue
        elif stripped.startswith('## '):
            add_h1(stripped[3:].strip())
        elif stripped.startswith('### '):
            add_h2(stripped[4:].strip())
        elif stripped.startswith('> '):
            quote_text = stripped[2:].strip()
            if quote_text.startswith('"') and quote_text.endswith('"'):
                quote_text = quote_text[1:-1]
            add_speech(quote_text)
        elif stripped.startswith('* '):
            bullet_text = stripped[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(4)
            if '**' in bullet_text:
                parts = bullet_text.split('**')
                for idx, part in enumerate(parts):
                    r = p.add_run(part)
                    r.font.color.rgb = RGBColor(0, 0, 0)
                    if idx % 2 == 1:
                        r.bold = True
            else:
                r = p.add_run(bullet_text)
                r.font.color.rgb = RGBColor(0, 0, 0)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run(stripped)
            r.font.color.rgb = RGBColor(0, 0, 0)

    output_path = os.path.join(os.path.dirname(__file__), '..', 'ExploQR-Development-Procedure-and-Presentation-Guide.docx')
    output_path = os.path.abspath(output_path)
    doc.save(output_path)
    print(f"Successfully generated plain Word document: {output_path}")

if __name__ == '__main__':
    create_document()
